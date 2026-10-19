"""知识库文档导入服务测试（不依赖 Qdrant / Ollama）

覆盖：各格式文本提取（txt/docx/pdf/xlsx/html）、语义切分、导入入库（mock upsert）、
不支持格式报错、空文本处理。

运行：conda base 环境下  python -m tests.kb_import_test
"""
import os
import sys

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND)

from io import BytesIO  # noqa: E402

import tyt.services.ai_service as ai  # noqa: E402
import tyt.services.kb_import_service as kb  # noqa: E402


# 用假 upsert 拦截，避免触碰 Qdrant；记录每次入库的 (doc_id, text, meta)
_calls = []


def _fake_upsert(doc_id, text, metadata=None):
    _calls.append((doc_id, text, metadata or {}))
    return True


ai.upsert_knowledge = _fake_upsert


def _make_docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_paragraph("产品风险排查操作手册")
    d.add_paragraph("第一章：入场前需核验工商登记信息与实控人穿透。")
    d.add_paragraph("第二章：关注司法诉讼与行政处罚记录的时效性。")
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_pdf() -> bytes:
    # 注意：reportlab 默认字体不支持中文，测试用 ASCII 文本验证提取链路；
    # 真实含内嵌中文字体的 PDF 可正常提取（pypdf 已验证）。
    from reportlab.pdfgen import canvas
    buf = BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "PDF Risk Notice Title")
    c.drawString(100, 730, "This notice describes the pre-investment risk self-check flow.")
    c.save()
    return buf.getvalue()


def _make_xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["指标", "数值"])
    ws.append(["资产负债率", "62%"])
    ws.append(["流动比率", "1.3"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_html() -> bytes:
    return ("<html><body><h1>HTML 帮助文档</h1>"
            "<p>这是 HTML 文档中的正文内容，用于验证网页导入。</p>"
            "<script>var x=1;</script></body></html>").encode("utf-8")


def test_extract_text_plain():
    t = kb.extract_text("a.txt", "标题行\n正文内容第一段。\n正文第二段。".encode("utf-8"))
    assert "标题行" in t and "正文内容第一段" in t


def test_extract_text_docx():
    t = kb.extract_text("手册.docx", _make_docx())
    assert "产品风险排查操作手册" in t and "实控人穿透" in t


def test_extract_text_pdf():
    t = kb.extract_text("notice.pdf", _make_pdf())
    assert "PDF Risk Notice Title" in t and "risk self-check flow" in t


def test_extract_text_xlsx():
    t = kb.extract_text("指标.xlsx", _make_xlsx())
    assert "资产负债率" in t and "流动比率" in t


def test_extract_text_html():
    t = kb.extract_text("帮助.html", _make_html())
    assert "HTML 帮助文档" in t and "网页导入" in t
    assert "var x=1" not in t  # script 应被剔除


def test_unsupported_format():
    try:
        kb.extract_text("a.exe", b"x")
        assert False, "应抛出 ValueError"
    except ValueError as e:
        assert "不支持" in str(e)


def test_chunk_text_split():
    long_text = "\n".join(f"第{i}段内容用于测试切分功能是否按段落正确拆分。" for i in range(50))
    chunks = kb.chunk_text(long_text, chunk_size=200, overlap=40)
    assert len(chunks) >= 2
    joined = "".join(chunks)
    # 所有原段落文字都应出现在某个分片中
    for i in range(0, 50, 7):
        assert f"第{i}段内容" in joined


def test_import_document_manual_kind():
    _calls.clear()
    res = kb.import_document("手册.docx", _make_docx())
    assert res["imported"] == res["chunks"] and res["imported"] > 0, res
    # 每条都标记为手动知识，且携带来源/分片元信息
    for doc_id, text, meta in _calls:
        assert meta["kind"] == "manual"
        assert meta["doc_name"] == "手册.docx"
        assert doc_id.startswith("manual:doc:")
        assert meta["chunk_total"] == res["chunks"]


def test_import_empty_text():
    res = kb.import_document("空.txt", b"   \n\n  ")
    assert res["imported"] == 0
    assert "未能提取" in res.get("message", "")


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items())
             if k.startswith("test_") and callable(v)]
    failed = 0
    for t in tests:
        _calls.clear()
        try:
            t()
            print(f"  PASS  {t.__name__}")
        except AssertionError as e:
            failed += 1
            print(f"  FAIL  {t.__name__}: {e}")
        except Exception as e:  # noqa
            failed += 1
            print(f"  ERROR {t.__name__}: {type(e).__name__}: {e}")
    print(f"\n结果：通过 {len(tests) - failed}/{len(tests)}")
    sys.exit(1 if failed else 0)
